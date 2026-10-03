"""File MCP: controlled access to the user's configured workspace.

Run: python -m mcp_servers.file_server   (stdio transport)
"""

import json
import mimetypes
import os
import re
import shutil
import time
from datetime import UTC, datetime
from pathlib import Path

from mcp.server.mcpserver import MCPServer

from mcp_servers.common import register
from mcp_servers.workspace import Workspace

ws = Workspace(os.environ.get("FILE_WORKSPACE", Path.home() / "AI-Agent-Workspace"))
mcp = MCPServer("file", instructions="Sandboxed access to the user's local workspace.")

MAX_READ_CHARS = 60_000
TEXT_EXTS = {".txt", ".md", ".markdown", ".py", ".js", ".ts", ".html", ".css", ".yaml", ".yml", ".log", ".xml"}


def _info(p: Path) -> dict:
    st = p.stat()
    return {
        "path": ws.rel(p),
        "type": "directory" if p.is_dir() else "file",
        "size": st.st_size,
        "modified": datetime.fromtimestamp(st.st_mtime, UTC).isoformat(),
        "mime_type": mimetypes.guess_type(p.name)[0],
    }


def _read_any(p: Path, max_rows: int = 200) -> str:
    """Single server, many formats: dispatch on extension."""
    ext = p.suffix.lower()
    if ext == ".json":
        return json.dumps(json.loads(p.read_text(encoding="utf-8")), indent=2)
    if ext in {".csv", ".tsv"}:
        import pandas as pd

        df = pd.read_csv(p, sep="\t" if ext == ".tsv" else ",")
        return f"{len(df)} rows × {len(df.columns)} columns\n" + df.head(max_rows).to_csv(index=False)
    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= max_rows:
                    rows.append("…")
                    break
                rows.append(",".join("" if v is None else str(v) for v in row))
            parts.append(f"## Sheet: {sheet.title}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    if ext == ".pdf":
        import fitz

        with fitz.open(p) as doc:
            return "\n\n".join(f"--- Page {i + 1} ---\n{page.get_text()}" for i, page in enumerate(doc))
    if ext == ".docx":
        from docx import Document

        return "\n".join(par.text for par in Document(p).paragraphs)
    if ext == ".pptx":
        from pptx import Presentation

        parts = []
        for i, slide in enumerate(Presentation(p).slides, start=1):
            texts = [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text.strip()]
            parts.append(f"--- Slide {i} ---\n" + "\n".join(texts))
        return "\n\n".join(parts)
    return p.read_text(encoding="utf-8", errors="replace")


def _clean_inline(t: str) -> str:
    """Strip Markdown emphasis markers so '**bold**' etc. don't show literally in Office docs."""
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"__(.+?)__", r"\1", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return t.strip()


def _parse_blocks(content: str):
    """Parse lightweight Markdown into (kind, level, text) tuples.

    kind ∈ {h1,h2,h3,bullet,number,text,blank}. Used to build structured Word/PowerPoint
    documents instead of dumping raw text."""
    out = []
    for raw in content.splitlines():
        indent = len(raw) - len(raw.lstrip(" "))
        s = raw.strip()
        if not s:
            out.append(("blank", 0, ""))
        elif s.startswith("### "):
            out.append(("h3", 0, _clean_inline(s[4:])))
        elif s.startswith("## "):
            out.append(("h2", 0, _clean_inline(s[3:])))
        elif s.startswith("# "):
            out.append(("h1", 0, _clean_inline(s[2:])))
        elif s[:2] in ("- ", "* ", "• "):
            out.append(("bullet", min(indent // 2, 4), _clean_inline(s[2:])))
        elif re.match(r"\d+\.\s+", s):
            out.append(("number", 0, _clean_inline(re.sub(r"^\d+\.\s+", "", s))))
        else:
            out.append(("text", 0, _clean_inline(s)))
    return out


def _write_docx(p: Path, content: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for kind, _level, text in _parse_blocks(content):
        if kind == "blank":
            continue
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "h3":
            doc.add_heading(text, level=3)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(text, style="List Number")
        else:
            doc.add_paragraph(text)
    doc.save(p)


def _write_pptx(p: Path, content: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9 widescreen
    blocks = [b for b in _parse_blocks(content) if b[0] != "blank"]
    deck_title = next((t for k, _, t in blocks if k == "h1"), "Presentation")

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = deck_title

    current_body = None
    current_heading = deck_title
    used_title = False

    def start_slide(heading: str):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
        slide.shapes.title.text = heading or deck_title
        tf = slide.placeholders[1].text_frame
        tf.clear()
        return tf

    for kind, level, text in blocks:
        if kind == "h1" and not used_title:
            used_title = True  # already used for the title slide
            continue
        if kind in ("h1", "h2"):
            current_heading = text
            current_body = start_slide(text)
            continue
        if not text or text == current_heading:
            continue  # skip blanks and a body line that just repeats the slide title
        if current_body is None:
            current_body = start_slide(deck_title)
        first = current_body.paragraphs[0]
        para = first if (first.text == "" and len(current_body.paragraphs) == 1) else current_body.add_paragraph()
        para.text = text
        para.level = 1 if kind == "h3" else (min(level, 4) if kind == "bullet" else 0)

    prs.save(p)


def _write_xlsx(p: Path, content: str) -> None:
    import csv
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    rows = list(csv.reader(io.StringIO(content)))
    for i, row in enumerate(rows):
        ws.append(row)
        if i == 0:  # bold header row
            for cell in ws[1]:
                cell.font = Font(bold=True)
    # Auto-size columns to content.
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    wb.save(p)


def _write_any(p: Path, content: str) -> None:
    ext = p.suffix.lower()
    p.parent.mkdir(parents=True, exist_ok=True)
    if ext == ".docx":
        _write_docx(p, content)
    elif ext == ".pptx":
        _write_pptx(p, content)
    elif ext in {".xlsx", ".xlsm"}:
        _write_xlsx(p, content)
    elif ext == ".json":
        p.write_text(json.dumps(json.loads(content), indent=2), encoding="utf-8")
    elif ext == ".pdf":
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 545, 792), content, fontsize=10)
        doc.save(p)
    else:
        p.write_text(content, encoding="utf-8")


@register(mcp, read_only=True)
def list_files(path: str = ".", recursive: bool = False) -> str:
    """List files and folders in a workspace directory."""
    d = ws.resolve(path, must_exist=True)
    it = d.rglob("*") if recursive else d.iterdir()
    items = [
        _info(p) for p in sorted(it) if not ws.rel(p).startswith(Workspace.TRASH_DIR) and not p.name.startswith(".")
    ][:500]
    return json.dumps({"directory": ws.rel(d), "items": items}, indent=1)


@register(mcp, read_only=True)
def search_files(query: str, path: str = ".", content: bool = False) -> str:
    """Find files whose name (or, with content=true, text content) contains the query."""
    d = ws.resolve(path, must_exist=True)
    q = query.lower()
    hits = []
    for p in d.rglob("*"):
        if not p.is_file() or ws.rel(p).startswith(Workspace.TRASH_DIR):
            continue
        if q in p.name.lower():
            hits.append({"path": ws.rel(p), "match": "name"})
        elif (
            content
            and p.suffix.lower() in TEXT_EXTS | {".csv", ".json"}
            and p.stat().st_size < 2_000_000
            and q in p.read_text(encoding="utf-8", errors="ignore").lower()
        ):
            hits.append({"path": ws.rel(p), "match": "content"})
        if len(hits) >= 100:
            break
    return json.dumps({"query": query, "results": hits})


@register(mcp, read_only=True)
def get_file_info(path: str) -> str:
    """Get metadata (size, type, modified time) for a file or folder. Returns exists=false if missing."""
    p = ws.resolve(path)
    if not p.exists():
        return json.dumps({"path": path, "exists": False})
    return json.dumps({"exists": True, **_info(p)})


@register(mcp, read_only=True)
def read_file(path: str) -> str:
    """Read a file. Supports txt, md, json, csv, xlsx, pdf, docx and other text files."""
    p = ws.resolve(path, must_exist=True)
    if p.is_dir():
        raise ValueError(f"'{path}' is a directory")
    text = _read_any(p)
    if len(text) > MAX_READ_CHARS:
        text = text[:MAX_READ_CHARS] + f"\n…[truncated, {len(text)} chars total]"
    return text


@register(mcp, destructive=True)
def write_file(path: str, content: str) -> str:
    """Write content to a file, creating it or overwriting an existing one. The file is built
    from `content` based on the extension — shape `content` accordingly:
    - .pptx (PowerPoint): Markdown. '# Deck Title' for the title slide, one '## Slide Title'
      per slide, '- ' bullet lines (two-space indent = sub-bullet, '### ' = a sub-point).
      Give every slide a '## ' heading.
    - .docx (Word): Markdown. '# '/'## '/'### ' headings, '- ' bullets, '1. ' numbered lists,
      blank lines between paragraphs.
    - .xlsx / .csv: CSV text, first row is the header.
    - .json: valid JSON. .pdf / .txt / .md / code: plain text.
    Always pass real content, never a placeholder like '[Your Name]'."""
    p = ws.resolve(path)
    existed = p.exists()
    _write_any(p, content)
    return json.dumps({"path": ws.rel(p), "overwritten": existed, "size": p.stat().st_size})


@register(mcp)
def create_file(path: str, content: str = "") -> str:
    """Create a new file (fails if it already exists). Use to make a Word doc (.docx),
    PowerPoint deck (.pptx), Excel sheet (.xlsx), PDF or text/code file. Shape `content` by
    extension:
    - .pptx: Markdown — '# Deck Title', then one '## Slide Title' per slide with '- ' bullets.
    - .docx: Markdown — '# '/'## '/'### ' headings, '- ' bullets, '1. ' numbered lists.
    - .xlsx / .csv: CSV text with a header row. .json: valid JSON. .pdf/.txt/.md: plain text.
    Always pass real content, never a placeholder."""
    p = ws.resolve(path)
    if p.exists():
        raise ValueError(f"'{path}' already exists; use write_file to overwrite")
    _write_any(p, content)
    return json.dumps({"path": ws.rel(p), "created": True, "size": p.stat().st_size})


@register(mcp)
def copy_file(source: str, destination: str) -> str:
    """Copy a file within the workspace."""
    src, dst = ws.resolve(source, must_exist=True), ws.resolve(destination)
    dst.parent.mkdir(parents=True, exist_ok=True)
    (shutil.copytree if src.is_dir() else shutil.copy2)(src, dst)
    return json.dumps({"copied": ws.rel(src), "to": ws.rel(dst)})


@register(mcp, destructive=True)
def move_file(source: str, destination: str) -> str:
    """Move a file or folder within the workspace."""
    src, dst = ws.resolve(source, must_exist=True), ws.resolve(destination)
    if dst.exists():
        raise ValueError(f"'{destination}' already exists")
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(src, dst)
    return json.dumps({"moved": ws.rel(src), "to": ws.rel(dst)})


@register(mcp, destructive=True)
def rename_file(path: str, new_name: str) -> str:
    """Rename a file or folder in place (new_name is a bare name, not a path)."""
    if "/" in new_name or "\\" in new_name or new_name in {".", ".."}:
        raise ValueError("new_name must be a plain file name")
    src = ws.resolve(path, must_exist=True)
    dst = ws.resolve(ws.rel(src.parent) + "/" + new_name)
    if dst.exists():
        raise ValueError(f"'{new_name}' already exists")
    src.rename(dst)
    return json.dumps({"renamed": path, "to": ws.rel(dst)})


@register(mcp, destructive=True)
def delete_file(path: str) -> str:
    """Delete a file or folder. It is moved to the workspace trash so it can be recovered."""
    p = ws.resolve(path, must_exist=True)
    if p == ws.root:
        raise ValueError("Refusing to delete the workspace root")
    trash = ws.root / Workspace.TRASH_DIR / f"{int(time.time())}_{p.name}"
    trash.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(p, trash)
    return json.dumps({"deleted": path, "recoverable_at": ws.rel(trash)})


if __name__ == "__main__":
    mcp.run()
