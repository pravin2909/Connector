"""Document parsers: turn a file into structural blocks (text + page + section).

Keeping page/section on every block is what makes citations possible later.
"""

import json
import re
import statistics
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_TYPES = {"pdf", "docx", "txt", "csv", "md", "json"}
_EXT_ALIASES = {"markdown": "md", "text": "txt"}


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None
    # Blocks that must not be merged with neighbours or split mid-row (e.g. CSV rows).
    atomic: bool = False


def file_type_of(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    return _EXT_ALIASES.get(ext, ext)


def parse(path: Path) -> tuple[list[Block], int | None]:
    """Returns (blocks, page_count)."""
    ft = file_type_of(path)
    if ft not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file type: .{ft}")
    return {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "md": _parse_markdown,
        "txt": _parse_txt,
        "csv": _parse_csv,
        "json": _parse_json,
    }[ft](path)


def _parse_pdf(path: Path):
    import fitz

    blocks: list[Block] = []
    with fitz.open(path) as doc:
        # Estimate body font size across the document; noticeably larger short lines are headings.
        sizes = [
            round(span["size"], 1)
            for page in doc
            for b in page.get_text("dict")["blocks"]
            for line in b.get("lines", [])
            for span in line["spans"]
            if span["text"].strip()
        ]
        body = statistics.median(sizes) if sizes else 11
        section = None
        for pno, page in enumerate(doc, start=1):
            buf: list[str] = []
            for b in page.get_text("dict")["blocks"]:
                for line in b.get("lines", []):
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    if not text:
                        continue
                    size = max(s["size"] for s in line["spans"])
                    bold = any(s["flags"] & 16 for s in line["spans"])
                    if len(text) < 120 and (size >= body * 1.18 or (bold and size >= body and len(text) < 80)):
                        if buf:
                            blocks.append(Block("\n".join(buf), pno, section))
                            buf = []
                        section = text
                    else:
                        buf.append(text)
                buf.append("")  # paragraph break between layout blocks
            if any(x.strip() for x in buf):
                blocks.append(Block("\n".join(buf).strip(), pno, section))
        return blocks, doc.page_count


def _parse_docx(path: Path):
    from docx import Document

    doc = Document(path)
    blocks: list[Block] = []
    section, buf = None, []
    for par in doc.paragraphs:
        style = (par.style.name or "").lower() if par.style is not None else ""
        if style.startswith("heading") or style == "title":
            if buf:
                blocks.append(Block("\n".join(buf), None, section))
                buf = []
            section = par.text.strip() or section
        elif par.text.strip():
            buf.append(par.text)
    if buf:
        blocks.append(Block("\n".join(buf), None, section))
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        blocks.append(Block("\n".join(rows), None, f"Table {t_idx}"))
    return blocks, None


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def _parse_markdown(path: Path):
    blocks: list[Block] = []
    trail: list[str] = []  # heading breadcrumb, e.g. "Architecture > Backend"
    buf: list[str] = []
    in_code = False
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else _MD_HEADING.match(line)
        if m:
            if any(x.strip() for x in buf):
                blocks.append(Block("\n".join(buf).strip(), None, " > ".join(trail) or None))
            buf = []
            level = len(m.group(1))
            trail = trail[: level - 1] + [m.group(2).strip()]
        else:
            buf.append(line)
    if any(x.strip() for x in buf):
        blocks.append(Block("\n".join(buf).strip(), None, " > ".join(trail) or None))
    return blocks, None


def _parse_txt(path: Path):
    return [Block(path.read_text(encoding="utf-8", errors="replace"))], None


def _parse_csv(path: Path, rows_per_block: int = 25):
    import pandas as pd

    df = pd.read_csv(path)
    header = ",".join(map(str, df.columns))
    blocks = []
    for start in range(0, len(df), rows_per_block):
        part = df.iloc[start : start + rows_per_block]
        # Repeat the header in every block so each chunk is self-describing.
        text = header + "\n" + part.to_csv(index=False, header=False)
        blocks.append(Block(text, None, f"Rows {start + 1}-{start + len(part)}", atomic=True))
    return blocks, None


def _parse_json(path: Path):
    data = json.loads(path.read_text(encoding="utf-8"))
    blocks = []
    if isinstance(data, dict):
        for key, value in data.items():
            blocks.append(Block(json.dumps(value, indent=1, ensure_ascii=False), None, str(key)))
    elif isinstance(data, list):
        for i, item in enumerate(data):
            blocks.append(Block(json.dumps(item, indent=1, ensure_ascii=False), None, f"Item {i + 1}"))
    else:
        blocks.append(Block(json.dumps(data)))
    return blocks, None
